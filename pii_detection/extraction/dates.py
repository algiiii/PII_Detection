"""Reference date of a document, with the provenance of the estimate (block B3).

The retention check (B7) asks how old a document is, and the file system alone is a
poor witness: copying a share, migrating it or restoring it from a backup rewrites
every ``mtime``, and a document from 2019 suddenly looks like it was written
yesterday. Born-digital files usually carry a better answer inside themselves — the
PDF ``modDate``, the Word or Excel core properties — so this module reads that first
and falls back to the ``mtime`` only when it has to.

Whatever it ends up using, it says so: a :class:`ReferenceDate` carries the
:class:`DateSource` it came from and the exact field it was read out of, so the
compliance verdict can qualify how much the signal is worth instead of presenting a
guess as a fact. The date is *semantic* — how old the content is — and must not be
confused with the technical "has this file changed since we last looked?" question,
which is answered by the file stamp of the registry.

Readers are imported lazily, like the text extractors: importing this module costs
nothing, only reading a file pulls PyMuPDF, python-docx or openpyxl.
"""

from __future__ import annotations

import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from enum import StrEnum
from pathlib import Path

#: Earliest date considered a plausible document date; anything older is taken as a
#: broken or default-initialised metadata field rather than a real date.
_MIN_PLAUSIBLE = datetime(1990, 1, 1, tzinfo=timezone.utc)

#: How far ahead of "now" a metadata date may sit before being rejected — enough to
#: absorb time-zone and clock skew, not enough to accept a date in the future.
_FUTURE_TOLERANCE = timedelta(days=1)

#: PDF date syntax (``D:YYYYMMDDHHmmSS+01'00'``). Everything after the day is
#: optional: real files, especially old ones, truncate freely.
_PDF_DATE = re.compile(
    r"D?:?"
    r"(?P<year>\d{4})(?P<month>\d{2})(?P<day>\d{2})"
    r"(?:(?P<hour>\d{2})(?:(?P<minute>\d{2})(?:(?P<second>\d{2}))?)?)?"
    r"\s*(?:(?P<tz>[Zz])|(?P<sign>[+-])(?P<tz_hour>\d{2})'?(?:(?P<tz_minute>\d{2})'?)?)?"
)


class DateSource(StrEnum):
    """Where a document's reference date was read from.

    :cvar CONTENT_METADATA: a date stored inside the file itself (PDF info
        dictionary, Word or Excel core properties) — the better estimate, since it
        survives copies and restores.
    :cvar FILE_MTIME: the file system's last-modified time — the fallback, and a
        weak signal: any bulk copy or backup restore resets it.
    """

    CONTENT_METADATA = "content_metadata"
    FILE_MTIME = "file_mtime"


@dataclass(frozen=True)
class ReferenceDate:
    """The date a document is assumed to date from, and how it was obtained.

    :ivar value: the reference date, always timezone-aware and in UTC.
    :ivar source: which kind of evidence produced it.
    :ivar field: the exact field it came from (``"pdf:modDate"``,
        ``"docx:core_properties.created"``, ``"fs:mtime"``, …), for the audit trail.
    """

    value: datetime
    source: DateSource
    field: str


def as_utc(value: datetime) -> datetime:
    """Read a naive datetime as UTC, leave an aware one alone.

    ``python-docx`` and ``openpyxl`` hand back naive datetimes, and so does SQLite
    when a stored timestamp is read back; coercing in one shared place means every
    comparison in the system mixes aware values only, and no caller has to
    remember the rule.

    :param value: the datetime to normalize.
    :returns: the equivalent timezone-aware datetime.
    """
    return value if value.tzinfo is not None else value.replace(tzinfo=timezone.utc)


def _parse_pdf_date(raw: str) -> datetime | None:
    """Parse a PDF date string into an aware datetime.

    Tolerant by design: the time, the seconds and the offset are all optional,
    because truncated values are common in the wild and a strict parser would return
    nothing exactly on the oldest files.

    :param raw: the raw value of a PDF info-dictionary date field.
    :returns: the parsed datetime in UTC, or ``None`` if it is not a PDF date.
    """
    match = _PDF_DATE.match(raw.strip())
    if match is None:
        return None
    parts = match.groupdict()
    if parts["sign"] is not None:
        offset = timedelta(
            hours=int(parts["tz_hour"]), minutes=int(parts["tz_minute"] or 0)
        )
        tzinfo = timezone(-offset if parts["sign"] == "-" else offset)
    else:
        tzinfo = timezone.utc
    try:
        value = datetime(
            int(parts["year"]),
            int(parts["month"]),
            int(parts["day"]),
            int(parts["hour"] or 0),
            int(parts["minute"] or 0),
            int(parts["second"] or 0),
            tzinfo=tzinfo,
        )
    except ValueError:  # syntactically fine, calendrically impossible (month 13, …)
        return None
    return value.astimezone(timezone.utc)


def _pdf_dates(path: Path) -> list[tuple[str, datetime]]:
    """Read the modification and creation dates of a PDF, in that order."""
    import fitz  # lazy: heavy dependency, only needed for PDFs

    with fitz.open(path) as document:
        metadata = document.metadata or {}
    candidates: list[tuple[str, datetime]] = []
    for key in ("modDate", "creationDate"):
        raw = metadata.get(key)
        parsed = _parse_pdf_date(raw) if raw else None
        if parsed is not None:
            candidates.append((f"pdf:{key}", parsed))
    return candidates


def _docx_dates(path: Path) -> list[tuple[str, datetime]]:
    """Read the core properties of a ``.docx``, modification date first."""
    from docx import Document  # lazy

    properties = Document(str(path)).core_properties
    return [
        (f"docx:core_properties.{name}", as_utc(value))
        for name, value in (
            ("modified", properties.modified),
            ("created", properties.created),
        )
        if value is not None
    ]


def _xlsx_dates(path: Path) -> list[tuple[str, datetime]]:
    """Read the workbook properties of a ``.xlsx``/``.xlsm``, modification first.

    Opened read-only: only the properties are wanted, not the sheets.
    """
    from openpyxl import load_workbook  # lazy

    workbook = load_workbook(path, read_only=True)
    try:
        properties = workbook.properties
        return [
            (f"xlsx:properties.{name}", as_utc(value))
            for name, value in (
                ("modified", properties.modified),
                ("created", properties.created),
            )
            if value is not None
        ]
    finally:
        workbook.close()


#: Formats whose files carry usable internal dates. Anything absent from this map —
#: ``.txt``, ``.doc``, ``.ods`` — goes straight to the ``mtime``, with no special
#: case anywhere else.
_METADATA_READERS: dict[str, Callable[[Path], list[tuple[str, datetime]]]] = {
    ".pdf": _pdf_dates,
    ".docx": _docx_dates,
    ".xlsx": _xlsx_dates,
    ".xlsm": _xlsx_dates,
}


def _is_plausible(value: datetime, now: datetime) -> bool:
    """Whether a metadata date can be believed.

    The single home of the sanity rule: dirty metadata is the norm (epoch defaults,
    clocks set wrong, dates in the future), and a guard spread across the format
    readers would drift between them.

    :param value: the candidate date, timezone-aware.
    :param now: the current time.
    :returns: ``True`` if the date is neither absurdly old nor in the future.
    """
    return _MIN_PLAUSIBLE <= value <= now + _FUTURE_TOLERANCE


def reference_date(path: str | Path, *, now: datetime | None = None) -> ReferenceDate:
    """Determine the date a document is assumed to date from.

    Prefers, in order: the modification date stored inside the file, its creation
    date, and finally the file system's ``mtime``. A metadata date that fails
    :func:`_is_plausible` is discarded in favour of the next candidate. Metadata that
    cannot be read at all — a corrupt PDF, unreadable properties — is not an error
    here: the ``mtime`` fallback is already the right answer, and refusing to date a
    document would abort its ingestion for no gain.

    :param path: path to the file to date; any format is accepted, including those
        with no metadata reader.
    :param now: current time, injectable so the "date in the future" guard is
        testable; defaults to the current UTC time.
    :returns: the :class:`ReferenceDate`, with the provenance of the estimate.
    :raises FileNotFoundError: if the file does not exist.
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"document not found: {path}")
    now = now if now is not None else datetime.now(timezone.utc)

    reader = _METADATA_READERS.get(path.suffix.lower())
    if reader is not None:
        try:
            candidates = reader(path)
        except Exception:  # noqa: BLE001 — unreadable metadata must not stop ingestion
            candidates = []
        for field, value in candidates:
            value = as_utc(value)
            if _is_plausible(value, now):
                return ReferenceDate(
                    value=value, source=DateSource.CONTENT_METADATA, field=field
                )

    mtime = datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc)
    return ReferenceDate(value=mtime, source=DateSource.FILE_MTIME, field="fs:mtime")


__all__ = ["DateSource", "ReferenceDate", "as_utc", "reference_date"]
